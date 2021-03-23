# -*- coding: utf-8 -*-
import docx
from docx import Document #匯入庫

path = "E:\\python_data\\1234.docx" #檔案路徑
document = Document(path) #讀入檔案
tables = document.tables #獲取檔案中的表格集
table = tables[0 ]#獲取檔案中的第一個表格
for i in range(1,len(table.rows)):#從表格第二行開始迴圈讀取表格資料
result = table.cell(i,0).text + "" +table.cell(i,1).text + table.cell(i,2).text + table.cell(i,3).text
#cell(i,0)表示第(i+1)行第1列資料，以此類推
print(result)
_table_list = []
for i, row in enumerate(table.rows):   # 讀每行
    row_content = []
    for cell in row.cells:  # 讀一行中的所有單元格
        c = cell.text
        if c not in row_content:
            row_content.append(c)
    # print(row_content)
    _table_list.append(row_content)


